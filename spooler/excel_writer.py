"""Выгрузка заявки в формате эталонного листа «Евгений (2)»."""
from __future__ import annotations

from dataclasses import dataclass

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .reference import Reference, guess_entry
from .spools import SpoolRow

HEADERS = [
    ("титул", "Автоматически"),
    ("№ изометрии", "Вручную"),
    ("№ линии", "Автоматически"),
    ("Ревизия по оспуленой схеме", "Автоматически"),
    ("Spool", "Вручную"),
    ("Ident", "Вручную"),
    ("Перевод", "Автоматически"),
    ("Наименование", "Автоматически"),
    ("DN (мм)", "Автоматически"),
    ("Dd (дюймы)", "Автоматически"),
    ("Ед.изм (шт,м)", "Автоматически"),
    ("Кол-во", "Вручную"),
    ("Заявка №", "Вручную"),
    ("Примечание (получено/не получено)", ""),
    ("Умная логика примечание (получено/не получено) с листа WL", ""),
    ("Примечание", ""),
    ("Вспомогательный столбец для фильтрации", ""),
]
WIDTHS = [8, 42, 40, 10, 9, 26, 12, 52, 12, 12, 13, 10, 12, 18, 22, 18, 12]
FONT = "Arial"


@dataclass
class Issue:
    iso: str
    spool: str
    what: str
    detail: str


def _spool_order(name: str) -> int:
    digits = "".join(ch for ch in name if ch.isdigit())
    return int(digits) if digits else 0


def write_workbook(path: str, rows: list[SpoolRow], reference: Reference,
                   issues: list[Issue]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявка"

    bold = Font(name=FONT, bold=True)
    plain = Font(name=FONT)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    manual_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, ((title, mode), width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        letter = openpyxl.utils.get_column_letter(col)
        ws.column_dimensions[letter].width = width
        top = ws.cell(row=1, column=col, value=mode or None)
        top.font = plain
        if mode == "Вручную":
            top.fill = manual_fill
        head = ws.cell(row=2, column=col, value=title)
        head.font = bold
        head.fill = head_fill
        head.alignment = wrap
    ws.freeze_panes = "A3"

    for row in sorted(rows, key=lambda r: (r.iso, _spool_order(r.spool), r.ident)):
        entry = reference.lookup(row.ident)
        if entry is None:
            entry = guess_entry(row.ident, row.description)
            issues.append(Issue(
                row.iso, row.spool, "нет в справочнике «База»",
                f"{row.ident} — "
                + (f"тип определён по описанию как «{entry.kind}», проверьте и "
                   f"добавьте позицию в «Базу»" if entry
                   else "тип определить не удалось, заполните вручную")))
        title = row.iso.split("-")[1] if len(row.iso.split("-")) > 1 else ""
        qty = int(row.qty) if float(row.qty).is_integer() else row.qty
        values = [
            title,
            row.iso,
            row.line,
            row.revision,
            row.spool,
            row.ident,
            entry.kind if entry else "",
            entry.name if entry else row.description,
            row.size.replace("X", "x"),
            reference.inches(row.size),
            row.unit,
            qty,
            None,
            None,
            None,
            None,
            _spool_order(row.spool),
        ]
        ws.append(values)

    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.font = plain

    sheet = wb.create_sheet("Расхождения")
    sheet.append(["№ изометрии", "Спул", "Что не так", "Подробности"])
    for cell in sheet[1]:
        cell.font = bold
        cell.fill = head_fill
    for width, letter in zip((42, 10, 44, 70), "ABCD"):
        sheet.column_dimensions[letter].width = width
    for issue in issues:
        sheet.append([issue.iso, issue.spool, issue.what, issue.detail])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = plain
            cell.alignment = wrap
    if len(issues) == 0:
        sheet.append(["", "", "расхождений не найдено", ""])

    wb.save(path)
